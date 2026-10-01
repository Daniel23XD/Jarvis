"""
JARVIS Actions Engine
Executes commands on the host laptop.
"""
import os
import urllib.parse
import webbrowser
from pathlib import Path
from docx import Document
from openpyxl import Workbook
import datetime

# User's Documents folder
USER_DIR = str(Path.home() / "Documents")

def execute_action(action_type: str, payload: str) -> dict:
    """Execute a system action based on LLM output."""
    action_type = action_type.strip().upper()
    payload = payload.strip()
    
    try:
        if action_type == "YOUTUBE":
            query = urllib.parse.quote(payload)
            webbrowser.open(f"https://www.youtube.com/results?search_query={query}")
            return {"status": "success", "message": f"Abriendo YouTube para buscar: {payload}"}
            
        elif action_type == "SPOTIFY":
            # For Windows, 'start spotify:search:...' opens the desktop app
            query = urllib.parse.quote(payload)
            os.system(f"start spotify:search:{query}")
            return {"status": "success", "message": f"Reproduciendo en Spotify: {payload}"}
            
        elif action_type == "SEARCH":
            query = urllib.parse.quote(payload)
            webbrowser.open(f"https://www.google.com/search?q={query}")
            return {"status": "success", "message": f"Buscando en la web: {payload}"}
            
        elif action_type == "WORD":
            doc = Document()
            doc.add_heading('Documento de JARVIS', 0)
            
            # Formatear el texto en párrafos para que se vea bien
            paragraphs = payload.split('\n')
            for p in paragraphs:
                if p.strip():
                    doc.add_paragraph(p.strip())
            
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"JARVIS_Doc_{timestamp}.docx"
            path = os.path.join(USER_DIR, filename)
            doc.save(path)
            
            # Open the file automatically
            os.system(f'start "" "{path}"')
            return {"status": "success", "message": f"Documento Word creado en Documentos: {filename}"}
            
        elif action_type == "EXCEL":
            wb = Workbook()
            ws = wb.active
            ws.title = "Datos JARVIS"
            # Simple parsing: comma separated payload
            rows = payload.split(";")
            for r_idx, row in enumerate(rows, 1):
                cols = row.split(",")
                for c_idx, val in enumerate(cols, 1):
                    ws.cell(row=r_idx, column=c_idx, value=val.strip())
            
            timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"JARVIS_Sheet_{timestamp}.xlsx"
            path = os.path.join(USER_DIR, filename)
            wb.save(path)
            
            # Open the file automatically
            os.system(f'start "" "{path}"')
            return {"status": "success", "message": f"Documento Excel creado en Documentos: {filename}"}
            
        return {"status": "error", "message": f"Acción desconocida: {action_type}"}
        
    except Exception as e:
        return {"status": "error", "message": f"Error ejecutando acción: {str(e)}"}
